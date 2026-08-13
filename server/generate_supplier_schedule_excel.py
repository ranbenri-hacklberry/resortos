"""
Supplier Schedule — Excel Generator & Importer
================================================
Generate:  python3 generate_supplier_schedule_excel.py <business_id>
Import:    python3 generate_supplier_schedule_excel.py <business_id> --import <file.xlsx>
"""
import sys, os, json, requests
from datetime import datetime

SUPABASE_URL = "http://127.0.0.1:54321"
SUPABASE_KEY = os.environ.get("SUPABASE_KEY", "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZS1kZW1vIiwicm9sZSI6InNlcnZpY2Vfcm9sZSIsImV4cCI6MTk4MzgxMjk5Nn0.EGIM96RAZx35lJzdJsyH-qQwv8Hdp7fsn3W0YpN81IU")
HEADERS = {"apikey": SUPABASE_KEY, "Authorization": f"Bearer {SUPABASE_KEY}", "Content-Type": "application/json", "Prefer": "return=representation"}

# Display: 1=ראשון to 6=שישי (no שבת) — DB stores 0-based (JS getDay())
DAY_NAMES_HE = ['ראשון', 'שני', 'שלישי', 'רביעי', 'חמישי', 'שישי']
DAY_MAP_HE = {v: i+1 for i, v in enumerate(DAY_NAMES_HE)}  # 1-based for display

def fetch_suppliers(business_id):
    url = f"{SUPABASE_URL}/rest/v1/suppliers?business_id=eq.{business_id}&select=id,name,delivery_schedule&order=name"
    r = requests.get(url, headers=HEADERS)
    r.raise_for_status()
    return r.json()

def generate_excel(business_id, output_path=None):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        print("Installing openpyxl...")
        os.system("pip install openpyxl")
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    suppliers = fetch_suppliers(business_id)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "לוח ספקים"
    ws.sheet_view.rightToLeft = True

    # Headers
    headers = ['ספק', 'יום אספקה', 'יום (1-6)', 'ימי הזמנה מראש', 'שעת סגירת הזמנה', 'הערות']
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name='Arial', bold=True, color="FFFFFF", size=11)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    # Data rows
    row = 2
    today_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
    tomorrow_fill = PatternFill(start_color="FFF8E1", end_color="FFF8E1", fill_type="solid")
    
    for supplier in suppliers:
        schedule = supplier.get('delivery_schedule') or []
        if not schedule:
            # Empty row for suppliers without schedule
            ws.cell(row=row, column=1, value=supplier['name']).font = Font(bold=True)
            ws.cell(row=row, column=2, value='—')
            ws.cell(row=row, column=3, value='')
            ws.cell(row=row, column=4, value=1)
            ws.cell(row=row, column=5, value='12:00')
            ws.cell(row=row, column=6, value='לא מוגדר')
            for col in range(1, 7):
                ws.cell(row=row, column=col).border = thin_border
            row += 1
            continue

        for entry in sorted(schedule, key=lambda x: x.get('day', 0)):
            day_db = entry.get('day', 0)  # 0-based from DB
            day_display = day_db + 1  # 1-based for display
            day_name = DAY_NAMES_HE[day_db] if 0 <= day_db <= 5 else '?'
            
            ws.cell(row=row, column=1, value=supplier['name']).font = Font(bold=True)
            ws.cell(row=row, column=2, value=day_name)
            ws.cell(row=row, column=3, value=day_display)
            ws.cell(row=row, column=4, value=entry.get('lead_days', 1))
            ws.cell(row=row, column=5, value=entry.get('cutoff', '12:00'))
            ws.cell(row=row, column=6, value=entry.get('notes', ''))
            
            for col in range(1, 7):
                ws.cell(row=row, column=col).border = thin_border
                ws.cell(row=row, column=col).alignment = Alignment(horizontal='center')
            ws.cell(row=row, column=1).alignment = Alignment(horizontal='right')
            
            row += 1

    # Column widths
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 16
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 20
    ws.column_dimensions['F'].width = 22

    # Instructions sheet
    ws2 = wb.create_sheet("הוראות")
    ws2.sheet_view.rightToLeft = True
    instructions = [
        "הוראות עריכה:",
        "",
        "ימי אספקה:",
        "  1 = ראשון",
        "  2 = שני",
        "  3 = שלישי",
        "  4 = רביעי",
        "  5 = חמישי",
        "  6 = שישי",
        "  (שבת — לא רלוונטי)",
        "",
        "עמודות:",
        "  ספק — שם הספק (חייב להתאים בדיוק לשם במערכת)",
        "  יום אספקה — ראשון/שני/שלישי/רביעי/חמישי/שישי",
        "  יום (1-6) — מספר היום (1=ראשון, 6=שישי)",
        "  ימי הזמנה מראש — 1=יום לפני, 2=יומיים לפני",
        "  שעת סגירת הזמנה — עד מתי אפשר להזמין (למשל 12:00)",
        "",
        "דוגמה:",
        "  ברכת האדמה | רביעי | 4 | 2 | 14:00",
        "  → אספקה ביום רביעי, הזמנה עד יום שני 14:00 (2 ימים לפני)",
        "",
        "להוספת יום אספקה חדש: הוסיפו שורה חדשה עם שם הספק",
        "למחיקת יום: מחקו את השורה",
        "",
        "אחרי עריכה, הריצו:",
        f"  python3 generate_supplier_schedule_excel.py {business_id} --import שם_הקובץ.xlsx"
    ]
    for i, line in enumerate(instructions, 1):
        ws2.cell(row=i, column=1, value=line).font = Font(name='Arial', size=11)
    ws2.column_dimensions['A'].width = 70

    # Save
    if not output_path:
        ts = datetime.now().strftime('%Y%m%d_%H%M')
        output_path = f"/Users/user/Documents/Max/supplier_schedule_{ts}.xlsx"
    
    wb.save(output_path)
    print(f"✅ Generated: {output_path}")
    print(f"   {len(suppliers)} suppliers, {row - 2} schedule entries")
    return output_path

def import_excel(business_id, file_path):
    try:
        import openpyxl
    except ImportError:
        os.system("pip install openpyxl")
        import openpyxl

    wb = openpyxl.load_workbook(file_path)
    ws = wb["לוח ספקים"]

    # Parse rows into schedule per supplier
    schedules = {}  # supplier_name -> [{ day, lead_days, cutoff }]
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        supplier_name = str(row[0] or '').strip()
        day_name = str(row[1] or '').strip()
        day_num = row[2]
        lead_days = row[3]
        cutoff = str(row[4] or '12:00').strip()
        notes = str(row[5] or '').strip()

        if not supplier_name or day_name == '—':
            continue

        # Resolve day number (Excel shows 1-based, DB stores 0-based)
        if day_num is not None and isinstance(day_num, (int, float)):
            day_num = int(day_num) - 1  # Convert 1-based display → 0-based DB
        elif day_name in DAY_MAP_HE:
            day_num = DAY_MAP_HE[day_name] - 1  # Convert 1-based → 0-based
        else:
            print(f"  ⚠️  Skipping {supplier_name}: unknown day '{day_name}'")
            continue

        if day_num < 0 or day_num > 5:
            print(f"  ⚠️  Skipping {supplier_name}: day {day_num} out of range (שבת לא רלוונטי)")
            continue

        if supplier_name not in schedules:
            schedules[supplier_name] = []
        
        entry = {"day": day_num, "lead_days": int(lead_days or 1), "cutoff": cutoff}
        if notes:
            entry["notes"] = notes
        schedules[supplier_name].append(entry)

    # Update each supplier
    updated = 0
    for name, schedule in schedules.items():
        url = f"{SUPABASE_URL}/rest/v1/suppliers?name=eq.{requests.utils.quote(name)}&business_id=eq.{business_id}"
        payload = {"delivery_schedule": schedule}
        r = requests.patch(url, headers=HEADERS, json=payload)
        if r.status_code in (200, 204):
            print(f"  ✅ {name}: {len(schedule)} ימי אספקה")
            updated += 1
        else:
            print(f"  ❌ {name}: {r.text}")

    # Clear schedule for suppliers not in Excel
    all_suppliers = fetch_suppliers(business_id)
    for s in all_suppliers:
        if s['name'] not in schedules and s.get('delivery_schedule'):
            url = f"{SUPABASE_URL}/rest/v1/suppliers?id=eq.{s['id']}"
            requests.patch(url, headers=HEADERS, json={"delivery_schedule": []})
            print(f"  🗑️  {s['name']}: cleared schedule")

    print(f"\n✅ Import complete! {updated} suppliers updated.")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage:")
        print("  Generate: python3 generate_supplier_schedule_excel.py <business_id>")
        print("  Import:   python3 generate_supplier_schedule_excel.py <business_id> --import <file.xlsx>")
        sys.exit(1)

    bid = sys.argv[1]

    if '--import' in sys.argv:
        idx = sys.argv.index('--import')
        if idx + 1 < len(sys.argv):
            import_excel(bid, sys.argv[idx + 1])
        else:
            print("❌ Missing file path after --import")
    else:
        generate_excel(bid)
